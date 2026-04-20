"""
Generate reports and outreach templates from scored opportunities.
Exports: Excel workbook (multi-sheet), CSV, and individual email templates.
"""
import os
import re
from datetime import datetime
import pandas as pd

EFFORT_EMOJI = {"low": "🟢", "medium": "🟡", "high": "🔴"}
PRIORITY_COLOR = {
    "A - Prioritaire": "FF4CAF50",  # green
    "B - Bon":         "FFFFC107",  # amber
    "C - Moyen":       "FFFF9800",  # orange
    "D - Faible":      "FFF44336",  # red
}

TEMPLATES = {
    "resource_page": """Objet : Suggestion de ressource pour votre page "{page_title}"

Bonjour,

En parcourant votre page {source_url}, j'ai remarqué que vous listez des ressources sur {topic}.

Je gère {your_site} qui propose {your_value_prop}. Je pense que cela pourrait intéresser vos lecteurs.

Seriez-vous ouvert à l'ajouter à votre liste ?

Cordialement,
{your_name}
""",

    "guest_post": """Objet : Proposition d'article invité – {topic}

Bonjour,

Je suis {your_name}, expert en {your_expertise}. J'ai beaucoup apprécié vos articles sur {source_url}.

Je souhaiterais vous proposer un article sur : "{article_idea}"

Voici quelques angles possibles :
- {angle_1}
- {angle_2}
- {angle_3}

Pensez-vous que cela correspondrait à votre ligne éditoriale ?

Cordialement,
{your_name}
""",

    "broken_link": """Objet : Lien cassé sur votre page {source_url}

Bonjour,

En visitant votre page {source_url}, j'ai remarqué qu'un lien pointe vers une page qui n'existe plus :
→ {broken_url}

Nous avons un contenu similaire qui pourrait le remplacer :
→ {your_url}

C'est une ressource sur {topic} qui est régulièrement mise à jour.

Je vous laisse en juger, bonne continuation !

Cordialement,
{your_name}
""",

    "testimonial": """Objet : Témoignage pour {tool_name}

Bonjour,

Voici mon témoignage sur {tool_name} que j'utilise via {your_site} :

"{testimonial_text}"

N'hésitez pas à l'utiliser sur votre site, idéalement avec un lien vers {your_site}.

Cordialement,
{your_name}
""",

    "directory": """— Soumission automatisable —
URL de soumission : {source_url}
Informations à préparer :
  - Nom du site : {your_site_name}
  - URL : {your_url}
  - Description (150 mots) : {your_description}
  - Catégorie : {category}
  - Email de contact : {your_email}
""",

    "unlinked_mention": """Objet : Mention de {your_brand} sur votre article

Bonjour,

J'ai remarqué que vous mentionnez {your_brand} dans votre article : {source_url}

Merci pour cette mention ! Seriez-vous ouvert à transformer cette mention en lien vers {your_url} ?
Cela permettrait à vos lecteurs d'accéder directement à la ressource.

Cordialement,
{your_name}
""",
}


def _col_letter(n: int) -> str:
    """Convert column index to Excel letter."""
    result = ""
    while n >= 0:
        result = chr(65 + n % 26) + result
        n = n // 26 - 1
    return result


def export_excel(df: pd.DataFrame, output_path: str, your_site: str = "") -> str:
    """Export scored opportunities to a formatted Excel workbook."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        return export_csv(df, output_path.replace(".xlsx", ".csv"))

    wb = Workbook()

    # ----- Sheet 1: All opportunities -----
    ws_all = wb.active
    ws_all.title = "Toutes les opportunités"
    _write_sheet(ws_all, df, your_site)

    # ----- Sheet 2: Priority A only -----
    priority_a = df[df["priority"] == "A - Prioritaire"].copy()
    if not priority_a.empty:
        ws_a = wb.create_sheet("Priorité A")
        _write_sheet(ws_a, priority_a, your_site)

    # ----- Sheet 3: By type summary -----
    ws_summary = wb.create_sheet("Résumé par type")
    summary = (
        df.groupby(["opportunity_label", "effort", "priority"])
          .agg(count=("rank", "count"), avg_score=("score", "mean"))
          .round(1)
          .reset_index()
          .sort_values("avg_score", ascending=False)
    )
    ws_summary.append(list(summary.columns))
    for _, row in summary.iterrows():
        ws_summary.append(list(row))

    wb.save(output_path)
    return output_path


def _write_sheet(ws, df: pd.DataFrame, your_site: str):
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter

    display_cols = [
        "rank", "priority", "score", "opportunity_label", "effort",
        "referring_domain", "source_url", "anchor_text",
        "authority_score", "dofollow", "strategy", "action"
    ]
    cols = [c for c in display_cols if c in df.columns]
    ws.append(cols)

    header_fill = PatternFill("solid", fgColor="FF1565C0")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.alignment = Alignment(horizontal="center")

    priority_fills = {
        "A - Prioritaire": PatternFill("solid", fgColor="FFE8F5E9"),
        "B - Bon":         PatternFill("solid", fgColor="FFFFF9C4"),
        "C - Moyen":       PatternFill("solid", fgColor="FFFFE0B2"),
        "D - Faible":      PatternFill("solid", fgColor="FFFFEBEE"),
    }

    for _, row in df[cols].iterrows():
        ws.append(list(row))
        prio = str(row.get("priority", ""))
        if prio in priority_fills:
            for cell in ws[ws.max_row]:
                cell.fill = priority_fills[prio]

    # Auto-width
    for col_idx, col in enumerate(cols, 1):
        max_len = max(
            df[col].astype(str).str.len().max() if col in df.columns else 0,
            len(col)
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)


def export_csv(df: pd.DataFrame, output_path: str) -> str:
    df.to_csv(output_path, index=False, encoding="utf-8-sig")
    return output_path


def generate_outreach_templates(df: pd.DataFrame, output_dir: str, your_site: str = "VOTRE_SITE"):
    """Generate one template file per opportunity type found."""
    os.makedirs(output_dir, exist_ok=True)
    generated = []
    for opp_type, template in TEMPLATES.items():
        subset = df[df["opportunity_type"] == opp_type]
        if subset.empty:
            continue
        path = os.path.join(output_dir, f"template_{opp_type}.txt")
        with open(path, "w", encoding="utf-8") as f:
            f.write(f"=== Template: {opp_type} ({len(subset)} opportunités) ===\n\n")
            f.write(template.replace("{your_site}", your_site))
            f.write("\n\n--- TOP 5 exemples ---\n")
            for _, row in subset.head(5).iterrows():
                f.write(f"  [{row.get('score', 0):.0f}pts] {row.get('source_url', '')}\n")
        generated.append(path)
    return generated


def print_summary(df: pd.DataFrame):
    """Print a colored summary to terminal."""
    print("\n" + "="*60)
    print(f"  RAPPORT NETLINKING — {len(df)} opportunités analysées")
    print("="*60)

    by_priority = df["priority"].value_counts().sort_index(ascending=False)
    for p, count in by_priority.items():
        print(f"  {p:<20} {count:>4} opportunités")

    print("\n  Top types d'opportunités :")
    by_type = df.groupby("opportunity_label")["score"].agg(["count", "mean"]).sort_values("mean", ascending=False)
    for label, row in by_type.head(8).iterrows():
        effort = df[df["opportunity_label"] == label]["effort"].iloc[0]
        emoji = EFFORT_EMOJI.get(effort, "⚪")
        print(f"  {emoji} {label:<35} {int(row['count']):>3}x  (score moy: {row['mean']:.0f})")

    print("\n  Top 10 opportunités :")
    top10 = df.nlargest(10, "score")[["rank", "score", "priority", "opportunity_label", "referring_domain"]]
    for _, r in top10.iterrows():
        print(f"  #{int(r['rank']):<3} [{r['score']:.0f}] {r['priority']:<20} {r['opportunity_label']:<30} {r['referring_domain']}")

    print("="*60 + "\n")
